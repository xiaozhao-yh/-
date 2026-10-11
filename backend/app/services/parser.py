"""简历文档解析：PDF / Word / TXT / HTML → 纯文本"""
import re

from fastapi import HTTPException


def _clean_text(text: str) -> str:
    text = text.lstrip("\ufeff")  # 去除 UTF-8 BOM
    lines = [ln.rstrip() for ln in text.splitlines()]
    out, blank = [], 0
    for ln in lines:
        if not ln.strip():
            blank += 1
            if blank > 1:
                continue
        else:
            blank = 0
        out.append(ln)
    return "\n".join(out).strip()


def _parse_pdf(data: bytes) -> str:
    import io

    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            t = page.extract_text() or ""
            if t.strip():
                pages.append(t)
        return "\n".join(pages)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"PDF 解析失败：{e}")


def _parse_docx(data: bytes) -> str:
    import io

    import docx as pydocx
    try:
        doc = pydocx.Document(io.BytesIO(data))
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Word 解析失败：{e}")


def _parse_txt(data: bytes) -> str:
    for enc in ("utf-8", "gbk", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_html(data: bytes) -> str:
    from bs4 import BeautifulSoup
    html = _parse_txt(data)
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text("\n")
    return text


def sniff_resume_file(filename: str, data: bytes) -> None:
    """上传文件安全检测：魔数校验（防扩展名伪造）+ Word 宏检测 + 二进制伪装检测"""
    import io
    import zipfile

    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        if not data[:8].startswith(b"%PDF-"):
            raise HTTPException(status_code=400, detail="文件内容与 PDF 格式不符（疑似伪造扩展名），已拒绝")
    elif ext == "docx":
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise HTTPException(status_code=400, detail="Word 文档格式异常或已损坏，已拒绝")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = zf.namelist()
        except zipfile.BadZipFile:
            raise HTTPException(status_code=400, detail="Word 文档格式异常或已损坏，已拒绝")
        if "[Content_Types].xml" not in names:
            raise HTTPException(status_code=400, detail="不是有效的 Word(.docx) 文档，已拒绝")
        for name in names:
            if "vba" in name.lower():  # vbaProject.bin / vbaData.xml 等宏部件
                raise HTTPException(status_code=400, detail="检测到文档包含 VBA 宏，为安全起见已拒绝，请另存为无宏的 .docx")
    elif ext in ("txt", "html", "htm"):
        if b"\x00" in data[:4096]:
            raise HTTPException(status_code=400, detail="检测到二进制内容伪装为文本文件，已拒绝")


def parse_resume_file(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        text = _parse_pdf(data)
    elif ext == "docx":
        text = _parse_docx(data)
    elif ext == "doc":
        raise HTTPException(status_code=400, detail="暂不支持旧版 .doc 格式，请另存为 .docx 后上传")
    elif ext == "txt":
        text = _parse_txt(data)
    elif ext in ("html", "htm"):
        text = _parse_html(data)
    else:
        raise HTTPException(status_code=400, detail="不支持的文件格式，请上传 PDF / Word / TXT / HTML")

    text = _clean_text(text)
    if not text:
        raise HTTPException(status_code=400, detail="未能从文件中解析出文本内容")
    return text


def normalize_text(text: str) -> str:
    return _clean_text(re.sub(r"\r\n", "\n", text))
