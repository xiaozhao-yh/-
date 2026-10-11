"""当前简历文本 → Word 文档导出"""
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BULLET_RE = re.compile(r"^[-•·*]\s+")


def _set_run_font(run, size: int | None = None, bold: bool = False, color: str | None = None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def export_resume_docx(title: str, text: str) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Pt(54)
        section.bottom_margin = Pt(54)
        section.left_margin = Pt(64)
        section.right_margin = Pt(64)

    # 文档标题
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(h.add_run(title or "个人简历"), size=18, bold=True)

    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            doc.add_paragraph()
            continue
        if line.lstrip().startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            content = line.lstrip("#").strip()
            p = doc.add_paragraph()
            _set_run_font(p.add_run(content), size=max(12, 16 - level), bold=True, color="1F4E79")
        elif BULLET_RE.match(line.strip()):
            content = BULLET_RE.sub("", line.strip())
            p = doc.add_paragraph(style="List Bullet")
            _set_run_font(p.add_run(content), size=11)
        else:
            p = doc.add_paragraph()
            _set_run_font(p.add_run(line), size=11)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
